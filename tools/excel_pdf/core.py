"""Read calculated Excel field columns and fill attorney-provided PDF templates.

This module never evaluates formulas, edits a workbook, submits a filing, or sends mail.
"""
from datetime import date, datetime
from pathlib import Path
import os
import re
import math

from openpyxl import load_workbook
from openpyxl.utils.escape import unescape
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject
from reportlab.pdfbase.pdfmetrics import stringWidth

FORMS = ("dc105", "dc508", "mc202")
EXCEL_ERRORS = {"#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
MIN_TEXT_SIZE = 7


class FieldValueError(ValueError):
    def __init__(self, field, message):
        super().__init__(message)
        self.field = field


def text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return unescape(value).replace("\r\n", "\n").replace("\r", "\n")
    if isinstance(value, (date, datetime)):
        if isinstance(value, datetime) and any((value.hour, value.minute, value.second, value.microsecond)):
            raise ValueError("Time-bearing cells need an explicit PDF display format")
        # The supplied completed PDFs use US dates with four-digit years.
        return f"{value.month}/{value.day}/{value.year}"
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError('Non-finite numeric value')
        return format(value, '.15g')
    if isinstance(value, (int, bool)):
        return str(value)
    raise ValueError('Expected a scalar calculated field value')


def row_values(workbook, row, sheet="Working Document"):
    if row < 2:
        raise ValueError("Choose a data row, not the header")
    cached = load_workbook(workbook, read_only=True, data_only=True)
    formulas = load_workbook(workbook, read_only=True, data_only=False)
    try:
        if sheet not in cached.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet}")
        source, raw = cached[sheet], formulas[sheet]
        if row > source.max_row:
            raise ValueError("Row is outside the worksheet")
        headers = next(source.iter_rows(min_row=1, max_row=1, values_only=True))
        cells = next(source.iter_rows(min_row=row, max_row=row))
        originals = next(raw.iter_rows(min_row=row, max_row=row))
        values, calculated, issues = {}, [], []
        for header, cell, original in zip(headers, cells, originals):
            if not header:
                continue
            name = str(header)
            if name in values:
                raise ValueError(f"Duplicate column header: {name}")
            values[name] = cell.value
            if original.data_type == "f" or hasattr(original.value, "text"):
                calculated.append(name)
                formula = getattr(original.value, 'text', original.value)
                if '#REF!' in str(formula):
                    issues.append({'field':name,'cell':cell.coordinate,'reason':'broken_formula_reference'})
                if cell.value is None and cell.data_type != 'str':
                    # Excel marks a cached empty string as t="str"; absent numeric caches are not safe.
                    issues.append({"field": name, "cell": cell.coordinate, "reason": "blank_formula_cache"})
            if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value in EXCEL_ERRORS):
                issues.append({"field": name, "cell": cell.coordinate, "reason": "excel_error", "value": cell.value})
        if not values.get("Defendant"):
            raise ValueError("The selected row has no Defendant")
        return {"sheet": sheet, "row": row, "values": values, "formula_fields": calculated, "issues": issues}
    finally:
        cached.close()
        formulas.close()


def page_widgets(page):
    for ref in page.get("/Annots", []):
        widget = ref.get_object()
        if widget.get("/Subtype") != "/Widget":
            continue
        field = widget
        visited = set()
        while "/T" not in field and field.get("/Parent"):
            if id(field) in visited:
                raise ValueError("Cyclic PDF field parent")
            visited.add(id(field))
            field = field["/Parent"].get_object()
        yield str(field.get("/T", "")), widget, field


def template_form(template, form):
    if form not in FORMS:
        raise ValueError("Choose dc105, dc508, or mc202")
    reader = PdfReader(template)
    if reader.is_encrypted:
        raise ValueError("Encrypted templates are not supported")
    pages = [i for i, page in enumerate(reader.pages)
             if any(name.startswith(form + "_") for name, _, _ in page_widgets(page))]
    if not pages or pages != list(range(min(pages), max(pages) + 1)):
        raise ValueError("The form must have one complete contiguous page range")
    writer = PdfWriter()
    writer.append(reader, pages=pages)
    original_form = reader.trailer['/Root']['/AcroForm']
    target_form = writer.root_object['/AcroForm']
    for name in ('/DR', '/DA'):
        if name in original_form:
            target_form[NameObject(name)] = original_form[name].clone(writer)
    resources = target_form.setdefault(NameObject('/DR'), DictionaryObject())
    fonts = resources.setdefault(NameObject('/Font'), DictionaryObject())
    # Adobe's supplied fields reference this standard font even though /DR omits it.
    if '/TiRo' not in fonts:
        fonts[NameObject('/TiRo')] = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Times-Roman'),
            NameObject('/Encoding'): NameObject('/WinAnsiEncoding')})
    fields = writer.get_fields() or {}
    if not fields or any(not name.startswith(form + "_") for name in fields):
        raise ValueError("Selected pages contain missing or unrelated canonical form fields")
    for page in writer.pages:
        for name, widget, field in page_widgets(page):
            if name not in fields:
                raise ValueError(f"Orphan PDF widget: {name}")
            if widget.get("/FT", field.get("/FT")) == "/Sig":
                raise ValueError("Digital-signature fields require a separate signing workflow")
            local_resources = widget.setdefault(NameObject('/DR'), DictionaryObject())
            local_fonts = local_resources.setdefault(NameObject('/Font'), DictionaryObject())
            for font, reference in fonts.items():
                if font not in local_fonts:
                    local_fonts[font] = reference
            appearance = str(widget.get('/DA', field.get('/DA', target_form.get('/DA', ''))))
            for font in re.findall(r'(/[^\s]+)\s+[\d.]+\s+Tf', appearance):
                if font not in local_fonts:
                    raise ValueError(f'Missing template font resource: {font}')
    return writer, fields, pages


def pdf_values(writer, fields, values):
    result = {}
    widgets = {name: (widget, field) for page in writer.pages for name, widget, field in page_widgets(page)}
    for name, field in fields.items():
        if name not in values:
            raise FieldValueError(name, f"Missing Excel column: {name}")
        value = values[name]
        if isinstance(value, str) and value in EXCEL_ERRORS:
            raise FieldValueError(name, f"Excel error in {name}: {value}")
        if field.get("/FT") == "/Btn":
            widget, _ = widgets[name]
            states = set(str(key) for key in widget.get("/AP", {}).get("/N", {})) - {"/Off"}
            if len(states) != 1:
                raise ValueError(f"Unsupported checkbox/radio states in {name}")
            if value in (None, "", False, "Off", "/Off", "No", "/No", "False", "FALSE"):
                result[name] = "/Off"
            elif value in (True, "On", "/On", "Yes", "/Yes", "True", "TRUE"):
                result[name] = next(iter(states))
            else:
                raise FieldValueError(name, f"Unrecognized checkbox value in {name}; use Yes/No or On/Off")
        else:
            if isinstance(value, bool):
                raise FieldValueError(name, f"Boolean value in text field {name}; confirm its formatting")
            try:
                result[name] = text(value)
            except ValueError as error:
                raise FieldValueError(name, f"Invalid value in {name}: {error}") from error
            try:
                result[name].encode('cp1252')
            except UnicodeEncodeError as error:
                raise FieldValueError(name, f'Template text font cannot represent the value in {name}') from error
            if any(ord(char) < 32 and char != '\n' or ord(char) == 127 for char in result[name]):
                raise FieldValueError(name, f'Unsupported control character in {name}')
    return result


def fit_text_widgets(writer, values):
    """Keep exact field values while preventing fixed-font appearance clipping."""
    acroform = writer.root_object['/AcroForm']
    for page in writer.pages:
        for name, widget, field in page_widgets(page):
            if widget.get('/FT', field.get('/FT')) != '/Tx' or not values[name]:
                continue
            value = values[name]
            multiline = bool(int(widget.get('/Ff', field.get('/Ff', 0))) & 4096)
            if '\n' in value and not multiline:
                raise FieldValueError(name, f'{name} is a single-line PDF field; remove line breaks in Excel')
            appearance = str(widget.get('/DA', field.get('/DA', acroform.get('/DA', ''))))
            fonts = list(re.finditer(r'(/[^\s]+)\s+([\d.]+)\s+Tf\b', appearance))
            if not fonts:
                raise ValueError(f'Missing template text font: {name}')
            font = fonts[-1]
            resource = widget['/DR']['/Font'][font[1]].get_object()
            family = str(resource.get('/BaseFont', '')).lstrip('/')
            if family not in ('Times-Roman', 'Helvetica', 'Courier'):
                raise ValueError(f'Unsupported template text font: {name}')
            size = float(font[2])
            if not math.isfinite(size) or size < MIN_TEXT_SIZE:
                raise ValueError(f'Unsupported template font size: {name}')
            rect = widget['/Rect']
            width = float(rect[2]) - float(rect[0]) - 4
            height = float(rect[3]) - float(rect[1]) - 2
            if width <= 0 or height <= 0:
                raise ValueError(f'Invalid template field dimensions: {name}')
            lines = value.split('\n')
            # Fixed-size pypdf appearances preserve explicit line breaks without wrapping.
            widest = max(stringWidth(line, family, size) for line in lines)
            fitted = min(size, size * width / widest) if widest else size
            if multiline:
                # Reserve full line boxes, including descenders on the last line.
                fitted = min(fitted, height / (1.2 * len(lines)))
            # Round down so PDF decimal precision cannot reintroduce clipping.
            fitted = math.floor(fitted * 100) / 100
            if fitted < MIN_TEXT_SIZE:
                raise FieldValueError(name, f'Text in {name} is too long to display legibly; '
                                      'shorten it in Excel' + (' or add line breaks' if multiline else ''))
            widget[NameObject('/DA')] = TextStringObject(
                appearance[:font.start(2)] + format(fitted, '.2f') + appearance[font.end(2):])


def generate(template, form, values, output):
    output = Path(output)
    if output.exists() or output.resolve() == Path(template).resolve():
        raise ValueError("Refusing to overwrite an existing file")
    writer, fields, pages = template_form(template, form)
    expected = pdf_values(writer, fields, values)
    fit_text_widgets(writer, expected)
    writer.update_page_form_field_values(None, expected, auto_regenerate=False, flatten=False)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(".pending.pdf")
    if temporary.exists():
        raise ValueError("A previous pending output exists; inspect it before retrying")
    try:
        with temporary.open("xb") as stream:
            writer.write(stream)
        reopened = PdfReader(temporary)
        actual = reopened.get_fields() or {}
        if set(actual) != set(expected):
            raise ValueError("Canonical field names changed during PDF export")
        for name, value in expected.items():
            if text(actual[name].get("/V")) != value:
                raise ValueError(f"Saved PDF value mismatch: {name}")
        count = 0
        for page in reopened.pages:
            for name, widget, field in page_widgets(page):
                count += 1
                if text(widget.get("/V", field.get("/V"))) != expected[name]:
                    raise ValueError(f"Widget value mismatch: {name}")
                if not widget.get("/AP", {}).get("/N"):
                    raise ValueError(f"Missing widget appearance: {name}")
                if widget.get('/FT', field.get('/FT')) == '/Btn':
                    # The supplied templates omit /Off from /N to render an unchecked box blank.
                    if str(widget.get('/AS')) != expected[name] or (expected[name] != '/Off' and expected[name] not in widget['/AP']['/N']):
                        raise ValueError(f'Checkbox appearance mismatch: {name}')
        if count < len(expected):
            raise ValueError("Export lost interactive widgets")
        os.link(temporary, output)
        temporary.unlink()
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return {"form": form, "pages": [n + 1 for n in pages], "field_count": len(expected),
            "widget_count": count, "fillable": True, "output": str(output)}
