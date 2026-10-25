"""Compare the supplied examples to exact calculated workbook columns, read-only."""
import argparse
import json
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from core import FORMS, row_values, template_form, text


def key(value):
    return ' '.join(text(value).split()).casefold()


def same_number(left, right):
    try:
        return Decimal(text(left)) == Decimal(text(right))
    except InvalidOperation:
        return False


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook', required=True)
    parser.add_argument('--template', required=True)
    parser.add_argument('--examples', nargs='+', required=True)
    parser.add_argument('--output-json', required=True)
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    source = workbook['Working Document']
    rows = source.iter_rows(values_only=True)
    headers = next(rows)
    positions = {name: index for index, name in enumerate(headers) if name}
    index = defaultdict(list)
    for number, row in enumerate(rows, 2):
        identity = tuple(key(row[positions[name]]) for name in ['dc105_01', 'dc105_05', 'dc105_06'])
        index[identity].append(number)
    workbook.close()
    template = PdfReader(args.template)
    template_fields = template.get_fields() or {}
    report = {'forms': {}, 'examples': []}
    for form in FORMS:
        _, fields, pages = template_form(args.template, form)
        report['forms'][form] = {'pages': [p + 1 for p in pages], 'field_count': len(fields),
                                'missing_columns': sorted(set(fields) - set(positions))}
    for example in args.examples:
        fields = PdfReader(example).get_fields() or {}
        identity = tuple(key(fields.get(name, {}).get('/V')) for name in ['dc105_01', 'dc105_05', 'dc105_06'])
        matches = index.get(identity, [])
        item = {'file': Path(example).name, 'matching_rows': matches, 'forms': {}}
        if len(matches) == 1:
            snapshot = row_values(args.workbook, matches[0])
            for form in FORMS:
                differences, count = [], 0
                for name, field in template_fields.items():
                    if not name.startswith(form + '_'):
                        continue
                    count += 1
                    source_value = snapshot['values'].get(name)
                    actual = fields.get(name, {}).get('/V')
                    if field.get('/FT') == '/Btn':
                        expected = source_value not in (None, '', False, 'Off', '/Off', 'No', '/No', 'False', 'FALSE')
                        same = expected == (str(actual or '/Off') not in ('/Off', '/No'))
                    else:
                        same = key(source_value) == key(actual)
                    if not same:
                        differences.append({'field': name, 'excel': text(source_value), 'pdf': text(actual),
                                            'numeric_equivalent': same_number(source_value, actual)})
                item['forms'][form] = {'compared_fields': count, 'differences': differences,
                    'cache_issues': [issue for issue in snapshot['issues'] if issue['field'].startswith(form + '_')]}
        report['examples'].append(item)
    output = Path(args.output_json)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps({'forms': report['forms'], 'examples': [
        {'file': item['file'], 'rows': item['matching_rows'], 'differences': {
            form: {'text_differences':len(result['differences']),
                   'material_differences':sum(not d['numeric_equivalent'] for d in result['differences'])}
            for form, result in item['forms'].items()}}
        for item in report['examples']]}, indent=2))


if __name__ == '__main__':
    main()
